"""只读核对原始SOP/体积表与交付XML；输出HTML/JSON，不改源文档、不操作GUI。"""
import argparse,ast,csv,hashlib,html,json,re
from pathlib import Path
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from docx import Document

BASE=Path('/Users/shan/Documents/BGI/SIRO解决方案中心/SIRO-Plus/实验')
SOP=BASE/'SOP/POCS呼吸道病原靶向单分子测序试剂盒（单分子纳米孔链测序法）产品说明书.docx'
XLSX=BASE/'TFT片上测试/片上测试1/TFT片上试剂体积表_单列.xlsx'
ACTUAL=BASE/'POCS/POCS正常体系预实验/SIRO-Plus_正常体系实验设计.xlsx'
SOFTWARE=BASE/'TFT片上测试/TFT编程相关/xPelabTFT软件需求说明书_260305更新.docx'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def rect(s):
    y,x,h,w=map(int,re.fullmatch(r'\((\d+),(\d+)\)\((\d+),(\d+)\)',s).groups());return x,y,w,h
def codes(r):return [rect(s) for s in r.get('RelayCode','').split(';') if s]

def audit(out):
    sheet=load_workbook(XLSX,data_only=True).worksheets[0]
    actual_book=load_workbook(ACTUAL,data_only=True)
    actual_names=' '.join(str(c.value) for s in actual_book for row in s for c in row if c.value is not None).replace('Ⅲ','III').replace(' ','')
    for name in ('PureScriptIIISupermix','VAHTSMulti-PCRBuffer','VAHTSRTMulti-PCREnzymeMix','En-Beads','RapidEndRepairReadyMix','RapidLigationBuffer','Ligase','EDTA'):
        assert name in actual_names,('实际试剂来源不匹配',name)
    software=Document(SOFTWARE)
    assert any('Temp on命令和delay命令' in p.text or 'Temo on命令和delay命令' in p.text for p in software.paragraphs)
    assert any('Accuracy' in c.text for t in software.tables for row in t.rows for c in row.cells)
    d=Document(SOP);paragraphs=[p.text for p in d.paragraphs]
    tables=[[[c.text for c in r.cells] for r in t.rows] for t in d.tables]
    table_text=lambda i:' '.join(' '.join(r) for r in tables[i])
    for i,expected in [(2,('37','10 min','85','5 s')),(3,('95','3 min','30 s','60 s','45 s','30 cycles','1 min')),
                       (5,('37','4 min')),(6,('2 μL','4 μL')),(7,('25','10 min','10','4'))]:
        assert all(x in table_text(i) for x in expected),('SOP表版本与审核基线不一致',i)
    assert any('室温静置5 min' in p and '磁力架3 min' in p for p in paragraphs)
    assert any('磁力架上静置2 min' in p and '15 µL' in p for p in paragraphs)
    assert any('磁力架上静置2 min' in p and '2.3.6.2' in p for p in paragraphs)
    assert any('磁力架上静置30 s' in p and '18 μL' in p for p in paragraphs)
    assert any('每2 min' in p for p in paragraphs)
    assert any('轻柔吹打20次' in p for p in paragraphs)
    root=ET.parse(out/'PTseq_16_A.elab').getroot();commands=list(root.iter('Command'))
    loads=list(csv.DictReader((out/'手工加液顺序.csv').open(encoding='utf-8-sig')))
    expected=[('待测核酸',[2]),('G4',[3]),('AMP',[6,7,8,9]),('G14',[11]),('G15',[12]),('G15',[13]),
              ('均一化产物',[14]),('G8',[15]),('G11 0.20 + G15 0.40 接头预混液',[17,18]),
              ('LIG',[19,20]),('EDTA',[21]),('G16',[23]),('G18',[24]),('G18',[25])]
    assert len(loads)==len(expected)
    rows=[]
    for item,(name,source_rows) in zip(loads,expected):
        planned=sum(float(sheet.cell(r,4).value) for r in source_rows)
        actual=float(item['本次µL（S/B孔为每道）'])/(16 if item['加液孔'].startswith('E') else 1)
        assert item['试剂']==name and abs(planned-actual)<1e-9,(name,planned,actual)
        if item['加液孔'].startswith('E'):
            ports=ast.literal_eval(item['各边缘孔µL'])
            assert abs(sum(ports.values())-actual*16)<1e-8
        for r in source_rows:
            assert abs(float(sheet.cell(r,4).value)/.05-int(sheet.cell(r,5).value))<1e-9
        rows.append(dict(reagent=' + '.join(str(sheet.cell(r,2).value) for r in source_rows),
                         excel_cells=','.join(f'D{r}' for r in source_rows),excel_ul=planned,program_ul=actual,
                         command=item['最后承接指令'],port=item['加液孔'],status='一致'))
    totals=[('A04',4,2),('A08',10,2.5),('R02-R03',None,5),('B04',16,2),('B09',None,4),('C02-C03',None,5.9)]
    for prefix,row_no,vol in totals:
        f=next(f for f in root if f.get('Name','').startswith(prefix))
        c=next(c for c in f if '循环至温控结束' in c.get('Name','') or '磁铁升起' in c.get('Name',''))
        occupied=set()
        for x,y,w,h in codes(c[-1]):occupied.update((xx,yy) for xx in range(x,x+w) for yy in range(y,y+h))
        for r in range(1,17):
            y0=27+9*(r-1);count=sum(10<=x<46 and y0<=y<y0+6 for x,y in occupied)
            assert abs(count*.05-vol)<1e-9,(prefix,r,count,vol)
        if row_no:assert abs(float(sheet.cell(row_no,4).value)-vol)<1e-9
    assert float(sheet['D11'].value)/float(sheet['D10'].value)==1
    assert abs(float(sheet['D23'].value)-round(float(sheet['D22'].value)*.4/.05)*.05)<1e-8
    assert abs(float(sheet['D16'].value)+sum(float(sheet.cell(r,4).value) for r in range(17,22))-float(sheet['D22'].value))<1e-8
    thermal=[]
    expected_thermal=[('A04',[(37,600),(85,5),(4,30)],1),('A08',[(95,180)],1),
                      ('A09',[(95,30),(60,60),(72,45)],30),('A10',[(72,60),(4,30)],1),
                      ('B04',[(37,240),(4,30)],1),('B09',[(25,600),(4,30)],1)]
    for prefix,seq,repeat in expected_thermal:
        f=next(f for f in root if f.get('Name','').startswith(prefix));cs=list(f)
        actual=[]
        for i,c in enumerate(cs):
            if c.get('Component')=='Heater1':
                assert c.get('Action')=='CMD_TEMP_ON' and c.get('Params').endswith(',1')
                delay=cs[i+1];assert delay.get('Component')=='Delay' and delay.get('PreTask')==c.get('ID')+',FS'
                actual.append((int(c.get('Params').split(',')[0]),int(delay.get('Timeout'))))
                if len(actual)>1:assert c.get('PreTask')==cs[i-1].get('ID')+',FS'
        assert actual==seq and int(f.get('Repeats'))==repeat
        hv=next(c for c in cs if '循环至温控结束' in c.get('Name',''))
        assert hv.get('Params')=='SCRIPTCYCLE_TRUE' and hv.get('ID') in f.get('SyncEndingCmdList').split(',')
        assert cs[-2].get('ID') in f.get('SyncEndingCmdList').split(',')
        thermal.append(dict(section=prefix,setpoints_and_holds=actual,cycles=repeat,status='设定值及顺序一致',
                            runtime_attainment_verified=False))
    capture=[]
    for prefix,secs in [('R02-R03',180),('C02-C03',120),('R05',120),('R06',120),('C05',30),('C06',30)]:
        f=next(f for f in root if f.get('Name','').startswith(prefix) and any('磁铁升起' in c.get('Name','') for c in f))
        c=next(c for c in f if '磁铁升起' in c.get('Name','') and c.get('Component')=='HighVoltage')
        assert sum(int(r.get('Duration')) for r in c)==secs*1000
        capture.append(dict(section=prefix,seconds=secs,command=c.get('ID'),source='POCS纯化/洗脱静置段落'))
    contour=json.loads((out/'全流程代码校验_20261010.json').read_text())
    events=contour['external_magnet_sequence']
    assert [e['target'] for e in events]==['down','up','down','up','down','up','down','down','up','down','up','down','up','down']
    for e in contour['bead_elution_incubation_events']:
        assert e['initial_mix_cycles']==12
        if e['stage']=='R1':
            assert e['incubation_seconds']==300
            assert [p['start_seconds'] for p in e['periodic_mixing']]==[120,240]
            assert all(p['end_seconds']<=300 for p in e['periodic_mixing'])
        f=next(f for f in root if f.get('Name','').startswith(('R' if e['stage']=='R1' else 'C')+('05' if e['round']==1 else '06'))
               and any('室温洗脱孵育' in c.get('Name','') for c in f))
        c=next(c for c in f if '室温洗脱孵育' in c.get('Name',''))
        assert sum(int(r.get('Duration')) for r in c)==e['incubation_seconds']*1000,'孵育内轻晃计时不守恒'
    assert any('EDTA终止后轻柔往复混匀20轮' in c.get('Name','') for c in commands)
    deviations=[
        '纯化1：SOP向25µL扩增产物加25µL溶解液及50µL磁珠；当前按已批准方案省溶解液、按扩增产物1:1加2.50µL En-Beads，不是严格10×原配方。',
        '纯化1省略两轮80%乙醇清洗；纯化2省略极速清洗液清洗，沿用已批准试验改动。',
        '两次洗脱：纯化1每次5.00µL、纯化2每次3.00µL，均不属于标准SOP的单次10×洗脱。',
        '单well纯化后下片汇聚，不按SOP先混样再进行纯化2；16道手工均一化后回片1.30µL/道。',
        '纯化2磁珠0.4×4.20=1.68µL，取整1.70µL，实际0.4048×；结合体系5.90µL，跨中线为已授权例外。',
        '16样本档表4为21.45ng，13µL对应1.65ng/µL；SOP正文浓度阈值写1.7。表格采用较保守1.65，片上1.30µL对应2.145ng，不是74.25ng上机投入量。',
        '实际试剂品牌和SOP G编号不完全相同，按已批准Excel配比；本检查不证明品牌间生化等效性。',
        '混匀采用片上固定形状往返，轮数增强不代表与SOP枪头吹打/震荡的效果已经等效。',
        '4℃在工程中写为设定+30s保持，不是无限期4℃ hold；手工停靠维持电极不等于维持温控。',
        '105/75℃热盖只作为原SOP条件写在名称中，工程没有热盖硬件控制指令。'
    ]
    unresolved=[
        '磁铁由外部控制，XML没有XaxisMotor指令，也没有到位反馈或升降手工确认门；需确定外部同步或增加人工停靠后才能称为操作放行版本。',
        '软件需求文件确认Temp on→Delay串行及循环块的表达方式；参数末尾1是Accuracy精度，不是到温完成标志。60s是到温动作超时，不应理解为额外固定保温；到温回报、超时阻断及芯片液温仍需实机确认。',
        'SOP磁吸要求液体澄清；固定静置时间不能替代澄清/磁珠完整捕获的实物确认。',
        '本轮只完成代码及来源核对；微缩反应效率、洗脱回收率、实际磁珠残留及最终pool浓度尚未由实测放行。'
    ]
    report=dict(status='offline_sop_and_excel_check_pass_operator_release_pending',source_sop=str(SOP),source_xlsx=str(XLSX),
                source_actual_reagents=str(ACTUAL),source_software_requirements=str(SOFTWARE),
                hashes={**{str(p):sha(p) for p in (SOP,XLSX,ACTUAL,SOFTWARE)},'PTseq_16_A.elab':sha(out/'PTseq_16_A.elab')},
                source_files_unchanged=True,unit_tests_passed=15,volume_checks=rows,thermal_checks=thermal,
                magnet_capture_checks=capture,magnet_sequence=events,approved_sop_deviations=deviations,
                unresolved_release_items=unresolved,ide_compilation_verified=False,hardware_verified=False,
                visual_check_performed=False,nominal_seconds=contour['after_nominal_seconds'],
                previous_nominal_seconds=contour['before_nominal_seconds'],
                generic_xml_checker_errors=0,generic_xml_checker_warnings=['main/PTseq未在IDE重新编译（2项）','大汇液框及获准5.90µL体系超过通用5.2µL阈值','50V保持/55V移液沿用已有方案'])
    (out/'SOP与10x体系审核_20261010.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    esc=lambda v:html.escape(str(v))
    table=''.join('<tr>'+''.join(f'<td>{esc(r[k])}</td>' for k in ('reagent','excel_cells','excel_ul','program_ul','port','command','status'))+'</tr>' for r in rows)
    temp=''.join(f'<tr><td>{esc(t["section"])}</td><td>{esc(t["setpoints_and_holds"])}</td><td>{t["cycles"]}</td><td>{esc(t["status"])}</td></tr>' for t in thermal)
    mag=''.join(f'<tr><td>{esc(t["section"])}</td><td>{t["seconds"]}</td><td>{esc(t["command"])}</td></tr>' for t in capture)
    paragraphs=lambda items:''.join(f'<p>{esc(s)}</p>' for s in items)
    text=f'''<!doctype html><html lang="zh"><meta charset="utf-8"><title>SOP与10×体系审核</title>
<style>body{{font:15px/1.65 system-ui;color:#111;background:white;max-width:1200px;margin:36px auto;padding:0 24px}}table{{border-collapse:collapse;width:100%;margin:12px 0}}th,td{{padding:9px;text-align:left;border-bottom:1px solid #ccc;vertical-align:top}}h1{{font-size:24px}}h2{{font-size:19px;margin-top:28px}}p{{margin:9px 0}}@media print{{body{{margin:0}}tr{{break-inside:avoid}}}}</style>
<h1>单列A · SOP与10×体系审核 · 2026-10-10</h1>
<p>14处加液体积、反应体系电极面积、温度设定及PCR30循环通过离线核对。操作放行仍需解决外部磁铁同步及实机到温确认。本报告不是微缩生化效果验证。</p>
<p>来源：<a href="{SOP.as_uri()}">POCS原始产品说明书</a>；<a href="{XLSX.as_uri()}">单列10×体积表</a>。源文件保持不变。</p>
<h2>逐处加液核对（µL/道）</h2><table><tr><th>实际试剂</th><th>Excel位置</th><th>表格</th><th>工程</th><th>入口</th><th>指令</th><th>结果</th></tr>{table}</table>
<p>台面AMP预混：0.50 Buffer + 0.30 Enzyme + 0.50引物 + 0.70溶解液 = 2.00µL/道；连接预混1.00 Buffer + 0.40 Ligase = 1.40；接头0.20 + 稀释液0.40 = 0.60；终止后4.20µL，磁珠后5.90µL。</p>
<h2>温度与循环</h2><p>括号内为（℃，保温秒数）。升降温时间不按10×缩短。热盖为SOP记录条件，非硬件控制。</p><table><tr><th>阶段</th><th>设定/保温</th><th>循环</th><th>结果</th></tr>{temp}</table>
<h2>磁吸与混匀</h2><p>结合往复纯化1为20轮、纯化2为12轮；四次洗脱重悬均为12轮。纯化1的300s洗脱孵育在第120/240s各轻晃2轮，轻晃耗时包括在300s内。EDTA后新增20轮轻柔短程往复。均是片上工程动作，不证明与震荡或吹打等效。</p><table><tr><th>阶段</th><th>磁吸秒数</th><th>指令</th></tr>{mag}</table>
<p>顺序：磁铁下降时结合混匀/孵育；全部液滴到右侧后升起捕珠；整个排液期间保持升起；洗脱液先覆盖磁珠再下降重悬；孵育后回右侧再升起；两次产物均吸走后下降并排珠。当前这些是外部控制时序，不是XML电机动作或到位确认。</p>
<h2>已批准的试验偏离与口径</h2>{paragraphs(deviations)}
<h2>尚未完成的操作放行条件</h2>{paragraphs(unresolved)}
<p>本次15项离线单元测试和全流程逐帧校验通过。没有visual check、IDE重新编译或仪器执行。标称时长{report['nominal_seconds']/3600:.2f}小时（含温控调度字段；非实机用时），此前{report['previous_nominal_seconds']/3600:.2f}小时。手工等待时间另计。</p>
</html>'''
    (out/'SOP与10x体系审核_20261010.html').write_text(text,encoding='utf-8')
    print(json.dumps(dict(status=report['status'],volume_checks=len(rows),thermal_sections=len(thermal),capture_checks=len(capture),
                          nominal_seconds=report['nominal_seconds'],unresolved=len(unresolved)),ensure_ascii=False))
    return report

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('out',type=Path);audit(ap.parse_args().out)
