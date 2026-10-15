"""重建并核验授权修订；输出JSON，不操作GUI或仪器。"""
import argparse,csv,hashlib,json,re,runpy,sys
from pathlib import Path
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
import contour,geom
from scene import Scene


def validate(out):
    previous={};stats=dict(frames=0,moving_liquid_frames=0,reshaping_liquid_frames=0,padding_expansion_frames=0)
    original=Scene._record
    def record(s,dur=geom.STEP_MS,*a,**kw):
        original(s,dur,*a,**kw);stats['frames']+=1;current={}
        for d in s.drops.values():
            if not d.rects():continue
            cells=d.target_cells();rc=contour.bounds(cells)
            assert len(contour.rectangles(cells))==1
            x,y,w,h=rc
            for c in range(2,7):
                tx=10+40*(c-1)
                assert not (x<tx+18 and tx<x+w and y<172 and 23<y+h),('经过其他列TEC',d.name,rc,c)
            old=previous.get(id(d))
            if old and old[0] is d and old[1]==d.vol and len(old[2])!=len(cells):
                assert getattr(d,'g14_transport_padding',False) and d.vol==.5
                assert len(old[2])==10 and len(cells)==12 and rc[2:]==(3,4) and dur>=4000
                stats['padding_expansion_frames']+=1
            if old and old[0] is d and old[1]==d.vol and len(old[2])==len(cells) and old[2]!=cells:
                before=old[3];dx,dy=x-before[0],y-before[1]
                if rc[2:]==before[2:] and cells==frozenset((xx+dx,yy+dy) for xx,yy in old[2]):
                    assert abs(dx)+abs(dy)<=1,('平移跳格',d.name,before,rc)
                    stats['moving_liquid_frames']+=1
                else:
                    assert dur>=4000,('形变未驻留4秒',d.name,before,rc,dur)
                    if rc[2:]!=tuple(reversed(before[2:])):
                        assert max(w/before[2],before[2]/w,h/before[3],before[3]/h)<=2.5,('过大形变',d.name,before,rc)
                    stats['reshaping_liquid_frames']+=1
            current[id(d)]=(d,d.vol,cells,rc)
        previous.clear();previous.update(current)
    Scene._record=record
    saved=sys.argv
    try:
        sys.argv=['build.py','--cols','1','--out',str(out)]
        result=runpy.run_path(str(Path(__file__).with_name('build.py')),run_name='__main__')
    finally:Scene._record=original;sys.argv=saved
    p=result['P'];scripts=result['scripts'];root=ET.parse(out/'PTseq_16_A.elab').getroot()
    assert not p.S.drops
    ids=[e.get('ID') for e in root.iter() if e.tag in ('Flow','Command')]
    assert len(ids)==len(set(ids))
    for f in root:
        local={c.get('ID') for c in f}
        assert set(filter(None,f.get('SyncEndingCmdList','').split(',')))<=local
        for c in f:
            pre=c.get('PreTask','').split(',')[0];assert pre=='-1' or pre in local
            if c.get('Component')!='HighVoltage':continue
            relays=list(c);assert relays and all(r.get('RelayCode') for r in relays)
            assert sum(int(r.get('Duration')) for r in relays)==1000*int(c.get('Timeout'))
            for r in relays:
                assert int(r.get('Duration'))!=750
                if float(r.get('Amp'))==0:assert f.get('Name').startswith('A09 ')
    pcr=next(f for f in root if f.get('Name').startswith('A09 '));assert pcr.get('Repeats')=='30'
    pulse=next(c for c in pcr if c.get('Component')=='HighVoltage')
    assert [(int(r.get('Duration')),int(r.get('Amp'))) for r in pulse]==[(1000,0),(1000,50)]
    thermal_expected={'A04':[(37,600),(85,5),(4,30)],'A08':[(95,180)],
        'A09':[(95,30),(60,60),(72,45)],'A10':[(72,60),(4,30)],'B04':[(37,240),(4,30)],'B09':[(25,600),(4,30)]}
    for prefix,expected in thermal_expected.items():
        f=next(f for f in root if f.get('Name').startswith(prefix+' '));cs=list(f);actual=[]
        for i,c in enumerate(cs):
            if c.get('Component')=='Heater1':actual.append((int(c.get('Params').split(',')[0]),int(cs[i+1].get('Timeout'))))
        assert actual==expected,(prefix,actual)
    sheet_path=Path('/Users/shan/Documents/BGI/SIRO解决方案中心/SIRO-Plus/实验/TFT片上测试/片上测试1/TFT片上试剂体积表_单列.xlsx')
    sheet=load_workbook(sheet_path,data_only=True,read_only=True).worksheets[0]
    expected={'G4':float(sheet['D3'].value),'AMP':sum(float(sheet.cell(r,4).value) for r in (6,7,8,9)),
        'G14':float(sheet['D11'].value),'G8':float(sheet['D15'].value),
        'LIG':sum(float(sheet.cell(r,4).value) for r in (19,20)),'EDTA':float(sheet['D21'].value),'G16':float(sheet['D23'].value)}
    finishes=[e for e in p.S.pipeline_events if e['kind']=='finish']
    before={'G4':1.6,'AMP':.5,'G14':2.5,'G8':1.3,'LIG':2.6,'EDTA':4.,'G16':4.2}
    for e in finishes:
        assert all(abs(v-before[e['reagent']]-expected[e['reagent']])<1e-7 for v in e['final_volumes'].values())
    loads=result['rows'];assert len(loads)==16
    assert [x['per_lane_ul'] for x in loads if x['hole'].startswith('S')]==[1.6,5.,5.,3.,1.3,3.,3.,3.]
    barcode=next(x for x in loads if x['hole'].startswith('B'))
    assert barcode['per_lane_ul']==1.2
    bcuts=[e for e in p.S.split_events if e['parent'].startswith('K')]
    assert len(bcuts)==16 and all(e['volumes']==[.6,.6] for e in bcuts)
    for e in p.S.split_events:
        assert min(e['volumes'])>=.4*sum(e['volumes'])-1e-8
        assert all(0<=b-a<=1 for a,b in zip(e['gaps'],e['gaps'][1:]))
    mix=[c for fs in p.segs.values() for f in fs for c in f.cmds if hasattr(c,'mix_timing')]
    assert mix and all(c.params=='SCRIPTCYCLE_TRUE' and c.relays[0].codes==c.relays[-1].codes for c in mix)
    assert [(e['stage'],e['mix_cycles']) for e in p.bead_cleanup_events]==[('R1',12),('C',2)]
    assert p.S.manual_aspirations[-1][1]['volume_ul']==48
    assert len(p.S.manual_aspirations[-1][1]['port_volumes'])==4
    assert all(e['hardware_control'] is False for e in p.magnet_events)
    padding=p.S.transport_padding_events
    added=[e for e in padding if e['kind']=='add'];retired=[e for e in padding if e['kind']=='retire']
    assert len(added)==32 and stats['padding_expansion_frames']==32
    assert all(e['actual_volume_ul']==.5 and e['source_cells']==10 and e['target_cells']==12
               and e['rect'][2:]==[3,4] and e['settle_ms']==4000 for e in added)
    assert sum(e['cells'] for e in retired)==64
    merges=[e for e in p.S.reaction_merge_events if e.get('mode')=='direct_nearby_fusion']
    assert merges and all(e['precontact_reshape'] is False for e in merges)
    a03=next(f for f in root if f.get('Name').startswith('A03 '))
    a03_loop=next(c for c in a03 if c.get('Params')=='SCRIPTCYCLE_TRUE')
    assert list(a03_loop)[-1].get('Duration')=='1000'
    native_sha=hashlib.sha256((out/'PTseq_16_A.elab').read_bytes()).hexdigest()
    report=dict(status='passed',validation='offline_only',visual_check=False,hardware_verified=False,
        workbook_sha256=hashlib.sha256(sheet_path.read_bytes()).hexdigest(),script_sha256=native_sha,
        geometry=dict(all_frames_full_rectangles=True,ordinary_hole_gap=1,
            all_systems_one_grid_hole_clearance=True,actual_590_ul=5.9,electrode_590_cells=120,
            reaction_safe_band_height=geom.LANE_H,
            preferred_resident_shapes={str(v):list(geom.lane_shape(v)) for v in (1.3,1.6,2.,2.5,2.6,3.,4.,4.2,5.,5.9)},
            no_foreign_TEC=True,no_overlap=True,translation_max_cells=1,reshape_settle_ms=4000),
        g14_transport_padding=dict(actual_volume_ul=.5,electrode_cells=12,shape=[3,4],
            added_packets=len(added),added_cells=64,retired_cells=64,
            liquid_volume_changed=False,additional_tips=0,events=padding),
        direct_fusion=dict(precontact_reshape=False,host_held_until_contact=True,events=merges),
        timing=dict(step_ms=1000,pcr_cycles=30,pcr_pulse_ms=[1000,1000],mixing_closed_loops=len(mix),
            scriptcycle='TRUE',runtime_cycle_end_behavior_verified=False,
            a03_native_last_duration_ms=1000,a03_preview_tail_hold_runtime_verified=False),
        amounts=expected,manual_loading_groups=len(loads),manual_aspiration_groups=len(p.S.manual_aspirations),
        thermal=thermal_expected,magnet_events=p.magnet_events,elution_events=p.elution_events,
        frames=stats,flows=len(root),commands=sum(len(f) for f in root),
        split_events=len(p.S.split_events),split_minimum_fraction=.4,
        waste_maximum_port_count=4,notes=['计划液量与承接面积分别核算：5.90µL体系沿用120格框；G14的0.50µL运输滴暂用12格3×4，融合后取消该两格余量。',
            '奇数列单列16道；磁铁升降为外部操作时序，不是电机闭环或执行验证。'])
    (out/'修订校验_20261015.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(dict(status='passed',**stats,script_sha256=native_sha),ensure_ascii=False))
    return report


if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('out',type=Path);args=ap.parse_args();validate(args.out)
